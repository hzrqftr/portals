"""
Export the ledger block (columns A-C) of "Family Fund.xlsx" to the CSV that
scripts/import-fund.mjs reads. docs/coinbox-spec.md §11.

    py -3.12 scripts/fund-xlsx-to-csv.py "Family Fund.xlsx" fund.csv

Needs openpyxl (`py -3.12 -m pip install --user openpyxl`). On Windows, use the
`py` launcher: a bare `python` may be the Microsoft Store placeholder.

Only A-C are exported. E-F are SUMIF totals and H-K is a hand-ticked "paid"
grid; both are DERIVED in Coinbox now, and the grid had already drifted from the
ledger it summarised, so importing it would import a disagreement.

Amounts are written as exact decimal strings via round(x, 2): the workbook
stores money as floats (its balance cell reads 3165.4100000000008), and the
importer turns the string into integer sen without touching a float again.
"""

import csv
import sys
from datetime import date, datetime

import openpyxl


def main(src: str, dst: str) -> None:
    ws = openpyxl.load_workbook(src, data_only=True).worksheets[0]
    header = [ws.cell(1, c).value for c in (1, 2, 3)]
    if header != ["Month", "Description", "Amount"]:
        sys.exit(f"Unexpected header in A1:C1: {header}")

    rows = 0
    with open(dst, "w", newline="", encoding="utf-8") as out:
        w = csv.writer(out)
        w.writerow(["line", "date", "description", "amount"])
        for r in range(2, ws.max_row + 1):
            when, what, amount = (ws.cell(r, c).value for c in (1, 2, 3))
            if when is None and what is None and amount is None:
                continue
            if not isinstance(when, (date, datetime)) or what is None or amount is None:
                sys.exit(f"Row {r} is incomplete: {when!r}, {what!r}, {amount!r}")
            day = when.date() if isinstance(when, datetime) else when
            # `line` is the workbook row number, so a problem report points at
            # the cell the owner would open.
            w.writerow([r, day.isoformat(), str(what).strip(), f"{round(float(amount), 2):.2f}"])
            rows += 1
    print(f"{rows} rows -> {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
